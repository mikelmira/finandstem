"""Populate the 'Images' sheet in aquascaping-catalogue-seed.xlsx.

Reads the species list from the Images sheet (already pre-populated with
Wikipedia + Commons category URLs by build_catalogue_v2.py), then for each
species fetches up to 5 images from Wikimedia and writes file title, direct
URL, description URL, license, author, credit, and attribution-required flag
back into the same sheet.

Usage:
  pip install openpyxl requests
  python populate_images.py [path/to/aquascaping-catalogue-seed.xlsx]

Output: same xlsx, modified in place. The Images sheet's empty columns get
filled in. Re-running is safe — rows that already have a file URL are skipped.

REQUIRES INTERNET ACCESS to en.wikipedia.org and commons.wikimedia.org.
"""
from __future__ import annotations

import sys
import time
from pathlib import Path
from urllib.parse import quote, unquote

import requests
from openpyxl import load_workbook

USER_AGENT = "FinAndStem/0.1 (contact: mikee@dsg.co.za)"  # CHANGE if needed
WIKI_API = "https://en.wikipedia.org/w/api.php"
COMMONS_API = "https://commons.wikimedia.org/w/api.php"
HEADERS = {"User-Agent": USER_AGENT}
SLEEP = 0.4
PER_SPECIES = 5


def get_json(url: str, params: dict) -> dict:
    r = requests.get(url, params=params, headers=HEADERS, timeout=20)
    r.raise_for_status()
    return r.json()


def wiki_lead_image(article_title: str) -> str | None:
    data = get_json(WIKI_API, {
        "action": "query", "format": "json",
        "prop": "pageimages", "piprop": "name",
        "titles": article_title,
    })
    for _pid, page in data.get("query", {}).get("pages", {}).items():
        name = page.get("pageimage")
        if name:
            return f"File:{name}"
    return None


def commons_category_files(category: str, limit: int = 20) -> list[str]:
    cat = category if category.startswith("Category:") else f"Category:{category}"
    data = get_json(COMMONS_API, {
        "action": "query", "format": "json",
        "list": "categorymembers",
        "cmtitle": cat, "cmtype": "file", "cmlimit": limit,
    })
    return [m["title"] for m in data.get("query", {}).get("categorymembers", [])]


def commons_file_meta(file_title: str) -> dict | None:
    data = get_json(COMMONS_API, {
        "action": "query", "format": "json",
        "prop": "imageinfo",
        "iiprop": "url|extmetadata|mime|size",
        "titles": file_title,
    })
    for _pid, page in data.get("query", {}).get("pages", {}).items():
        infos = page.get("imageinfo")
        if not infos:
            continue
        info = infos[0]
        meta = info.get("extmetadata", {})

        def gv(k):
            v = meta.get(k)
            return v.get("value") if v else None

        license_short = gv("LicenseShortName") or ""
        # Skip non-free / fair use
        if any(b in license_short.lower() for b in ("fair", "non-free", "copyrighted", "all rights reserved")):
            return None
        if not (info.get("mime") or "").startswith("image/"):
            return None
        return {
            "fileTitle": file_title,
            "imageUrl": info.get("url"),
            "descriptionUrl": info.get("descriptionurl"),
            "license": license_short,
            "licenseUrl": gv("LicenseUrl"),
            "author": gv("Artist"),
            "credit": gv("Credit"),
            "attributionRequired": gv("AttributionRequired"),
            "mime": info.get("mime"),
            "width": info.get("width"),
            "height": info.get("height"),
        }
    return None


def wiki_title_from_url(url: str) -> str:
    return unquote(url.rstrip("/").split("/wiki/")[-1])


def commons_category_from_url(url: str) -> str:
    return unquote(url.rstrip("/").split("/wiki/")[-1])


def fetch_images_for_species(wiki_url: str, commons_cat_url: str) -> list[dict]:
    """Return up to PER_SPECIES image-metadata dicts."""
    results: list[dict] = []
    seen: set[str] = set()

    # 1) Wikipedia lead image
    try:
        title = wiki_title_from_url(wiki_url)
        lead = wiki_lead_image(title)
        time.sleep(SLEEP)
        if lead:
            meta = commons_file_meta(lead)
            time.sleep(SLEEP)
            if meta:
                seen.add(meta["fileTitle"])
                results.append(meta)
    except Exception as e:
        print(f"  ! lead image error: {e}", file=sys.stderr)

    # 2) Commons category
    try:
        cat = commons_category_from_url(commons_cat_url)
        files = commons_category_files(cat, limit=25)
        time.sleep(SLEEP)
        for ft in files:
            if len(results) >= PER_SPECIES:
                break
            if ft in seen:
                continue
            meta = commons_file_meta(ft)
            time.sleep(SLEEP)
            if not meta:
                continue
            seen.add(ft)
            results.append(meta)
    except Exception as e:
        print(f"  ! commons category error: {e}", file=sys.stderr)

    return results[:PER_SPECIES]


def main(xlsx_path: str) -> int:
    p = Path(xlsx_path)
    if not p.exists():
        print(f"File not found: {p}", file=sys.stderr)
        return 1

    wb = load_workbook(p)
    if "Images" not in wb.sheetnames:
        print("Sheet 'Images' not found in xlsx", file=sys.stderr)
        return 1
    ws = wb["Images"]

    # Map headers to column indexes (1-based for openpyxl)
    hdrs = {c.value: c.column for c in ws[1]}

    needed = [
        "Species ID", "Common Name", "Image #",
        "Wikipedia URL (article)", "Commons Category URL (gallery)",
        "Commons File Title", "Direct Image URL", "Description Page URL",
        "License (short)", "License URL", "Author / Photographer",
        "Credit", "Attribution Required",
    ]
    for n in needed:
        if n not in hdrs:
            print(f"Missing column: {n}", file=sys.stderr)
            return 1

    # Group rows by species ID
    species_rows: dict[str, list[int]] = {}
    species_info: dict[str, dict] = {}
    for r in range(2, ws.max_row + 1):
        sp_id = ws.cell(row=r, column=hdrs["Species ID"]).value
        if not sp_id:
            continue
        species_rows.setdefault(sp_id, []).append(r)
        if sp_id not in species_info:
            species_info[sp_id] = {
                "common": ws.cell(row=r, column=hdrs["Common Name"]).value,
                "wiki": ws.cell(row=r, column=hdrs["Wikipedia URL (article)"]).value,
                "cat": ws.cell(row=r, column=hdrs["Commons Category URL (gallery)"]).value,
            }

    total = len(species_info)
    for i, (sp_id, rows) in enumerate(species_rows.items(), 1):
        info = species_info[sp_id]
        print(f"[{i}/{total}] {sp_id} :: {info['common']}")

        # Check if first row already has a URL — skip if so
        first_row = rows[0]
        if ws.cell(row=first_row, column=hdrs["Direct Image URL"]).value:
            print("  (already populated, skipping)")
            continue

        try:
            metas = fetch_images_for_species(info["wiki"], info["cat"])
        except Exception as e:
            print(f"  ! fatal: {e}", file=sys.stderr)
            continue

        for idx, meta in enumerate(metas):
            if idx >= len(rows):
                break
            row = rows[idx]
            ws.cell(row=row, column=hdrs["Commons File Title"]).value = meta.get("fileTitle")
            ws.cell(row=row, column=hdrs["Direct Image URL"]).value = meta.get("imageUrl")
            ws.cell(row=row, column=hdrs["Description Page URL"]).value = meta.get("descriptionUrl")
            ws.cell(row=row, column=hdrs["License (short)"]).value = meta.get("license")
            ws.cell(row=row, column=hdrs["License URL"]).value = meta.get("licenseUrl")
            ws.cell(row=row, column=hdrs["Author / Photographer"]).value = meta.get("author")
            ws.cell(row=row, column=hdrs["Credit"]).value = meta.get("credit")
            ws.cell(row=row, column=hdrs["Attribution Required"]).value = str(meta.get("attributionRequired", "")).lower() == "true"

        # Save after each species so a crash doesn't lose progress
        wb.save(p)
        print(f"  -> {len(metas)} images saved")

    print("Done.")
    return 0


if __name__ == "__main__":
    path = sys.argv[1] if len(sys.argv) > 1 else "aquascaping-catalogue-seed.xlsx"
    sys.exit(main(path))
